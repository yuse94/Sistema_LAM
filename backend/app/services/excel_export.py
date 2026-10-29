"""Exportación a Excel con el formato de DB_LAM.xlsx (porta add_db_LAM.py).

La base de datos es la fuente de verdad; el libro se genera bajo demanda con una
hoja por vista (Anterior, Posterior, LateralD y LateralI) y enlaces al reporte PDF
y a la imagen de cada evaluación.

Cambios respecto a la versión de consola:
- Cada medición tiene dos columnas, dirección y valor (antes solo el valor, sin lado).
- El diagnóstico une clasificación y ángulo ("Valgo (3.2°)"); la versión de consola
  tomaba índices que dejaban solo el ángulo en las columnas de cabeza y rodillas.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .labels import MODE_LABELS, VIEW_SHEETS, VIEWS, diagnostic_lines, table_unit


@dataclass(frozen=True)
class TableLayout:
    title: str
    columns: tuple[str, str, str]
    items: tuple[str, ...]


@dataclass(frozen=True)
class SheetLayout:
    """Estructura de columnas de una vista (sale de evaluar la silueta de referencia)."""

    tables: tuple[TableLayout, ...]
    diagnostic_keys: tuple[str, ...]


@dataclass
class ExcelRow:
    view: str
    evaluation_id: int
    date: datetime
    mode: str
    patient: dict  # name, age, gender, weight_kg, occupation
    diagnostic: dict
    tables: list[dict]  # {title, columns, rows: [{item, direction, value}]}
    height_cm: float | None = None
    report_url: str | None = None
    image_url: str | None = None
    extra: dict = field(default_factory=dict)


HEADER_FILL = PatternFill("solid", fgColor="0E7C74")
SUBHEADER_FILL = PatternFill("solid", fgColor="E6F2F1")
THIN = Side(style="thin", color="B7C4C3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

BASE_GROUPS = (
    ("Enlaces", ("Reporte", "Imagen")),
    ("Evaluación", ("N.º", "Fecha", "Detección")),
    ("Paciente", ("Nombre", "Edad", "Género", "Talla [cm]", "Peso [kg]", "Ocupación")),
)


def _header(ws, row: int, col: int, text: str, *, rows: int = 1, cols: int = 1, level: int = 0) -> None:
    if rows > 1 or cols > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row + rows - 1, end_column=col + cols - 1)
    cell = ws.cell(row=row, column=col, value=text)
    cell.alignment = CENTER
    if level == 0:
        cell.fill, cell.font = HEADER_FILL, Font(bold=True, color="FFFFFF")
    else:
        cell.fill, cell.font = SUBHEADER_FILL, Font(bold=True, color="1F2933")
    for r in range(row, row + rows):
        for c in range(col, col + cols):
            ws.cell(row=r, column=c).border = BORDER


def _diagnostic_columns(keys: tuple[str, ...]) -> list[tuple[str, str]]:
    """(clave, etiqueta) sin las claves *_angulo, que van unidas a su clasificación."""
    return [(k, label) for k, label, _ in diagnostic_lines({k: "" for k in keys})]


def _write_sheet(ws, layout: SheetLayout, rows: list[ExcelRow]) -> None:
    col = 1
    for group, names in BASE_GROUPS:
        _header(ws, 1, col, group, cols=len(names))
        for i, name in enumerate(names):
            _header(ws, 2, col + i, name, rows=2, level=1)
        col += len(names)

    diag_cols = _diagnostic_columns(layout.diagnostic_keys)
    _header(ws, 1, col, "Diagnóstico", cols=len(diag_cols))
    for i, (_, label) in enumerate(diag_cols):
        _header(ws, 2, col + i, label, rows=2, level=1)
    col += len(diag_cols)

    first_table_col = col
    for t in layout.tables:
        unit = table_unit(t.columns)
        _header(ws, 1, col, t.title, cols=2 * len(t.items))
        for i, item in enumerate(t.items):
            _header(ws, 2, col + 2 * i, item, cols=2, level=1)
            _header(ws, 3, col + 2 * i, "Dir.", level=1)
            _header(ws, 3, col + 2 * i + 1, unit, level=1)
        col += 2 * len(t.items)
    last_col = col - 1

    for r, row in enumerate(rows, start=4):
        p = row.patient
        diag = {k: text for k, _, text in diagnostic_lines(row.diagnostic)}
        values = [
            None, None, row.evaluation_id, row.date.replace(tzinfo=None), MODE_LABELS.get(row.mode, row.mode),
            p.get("name"), p.get("age"), p.get("gender"), row.height_cm or p.get("height_cm"), p.get("weight_kg"),
            p.get("occupation"),
        ] + [diag.get(k) for k, _ in diag_cols]
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)
        for c, (url, text) in enumerate(((row.report_url, "PDF"), (row.image_url, "Imagen")), start=1):
            if url:
                cell = ws.cell(row=r, column=c, value=text)
                cell.hyperlink = url
                cell.font = Font(color="0E7C74", underline="single")
        ws.cell(row=r, column=4).number_format = "dd/mm/yyyy hh:mm"

        c = first_table_col
        for ti, t in enumerate(layout.tables):
            data_rows = row.tables[ti]["rows"] if ti < len(row.tables) else []
            for i in range(len(t.items)):
                if i < len(data_rows):
                    ws.cell(row=r, column=c, value=data_rows[i]["direction"])
                    cell = ws.cell(row=r, column=c + 1, value=round(float(data_rows[i]["value"]), 2))
                    cell.number_format = "0.00"
                c += 2

    widths = [8, 9, 6, 16, 18, 26, 6, 7, 9, 9, 18] + [16] * len(diag_cols)
    for i in range(1, last_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1] if i <= len(widths) else 8
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = ws.cell(row=4, column=7)  # fijos: encabezados y columnas hasta el nombre
    if rows:
        ws.auto_filter.ref = f"A3:{get_column_letter(last_col)}{3 + len(rows)}"


def build_workbook(rows: list[ExcelRow], layouts: dict[str, SheetLayout]) -> bytes:
    """Libro con una hoja por vista, aunque alguna no tenga evaluaciones."""
    wb = Workbook()
    wb.remove(wb.active)
    for view in VIEWS:
        ws = wb.create_sheet(VIEW_SHEETS[view])
        view_rows = sorted((r for r in rows if r.view == view), key=lambda r: (r.date, r.evaluation_id))
        _write_sheet(ws, layouts[view], view_rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
