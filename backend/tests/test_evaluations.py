"""Guardar evaluaciones, reporte PDF, Excel, historial y eliminación (LOPDP)."""
import base64
import io
import json

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app.core.config import get_settings
from app.main import app
from tests.synthetic import ANTERIOR, LATERAL, make_image
from tests.synthetic_pose import frontal_pose


def _png(arr) -> bytes:
    buf = io.BytesIO()
    Image.fromarray(arr).save(buf, format="PNG")
    return buf.getvalue()


def _analyze(c, view, points):
    """Detección con marcadores: devuelve la vista lista para guardar y la imagen enderezada."""
    r = c.post(f"/api/analysis/{view}", files={"file": ("f.png", _png(make_image(points)), "image/png")})
    assert r.status_code == 200, r.text
    body = r.json()
    image = base64.b64decode(body["image"].split(",", 1)[1])
    return {"view": view, "landmarks": body["markers"], "scale_ratio": body["scale_ratio"]}, image


def _save(c, payload, images):
    files = [("images", (f"{i}.jpg", img, "image/jpeg")) for i, img in enumerate(images)]
    return c.post("/api/evaluations", data={"data": json.dumps(payload)}, files=files)


def _patient(c, name="Ana Torres"):
    r = c.post("/api/patients", json={"name": name, "age": 30, "gender": "F", "weight_kg": 60})
    assert r.status_code == 201
    return r.json()["id"]


def test_save_evaluation_with_markers_and_outputs():
    with TestClient(app) as c:
        pid = _patient(c)
        ant, ant_img = _analyze(c, "anterior", ANTERIOR)
        lat, lat_img = _analyze(c, "lateral_derecha", LATERAL)
        payload = {"patient_id": pid, "mode": "markers", "notes": "  Control inicial ",
                   "config": {"tolerance_angle": 1, "tolerance_distance": 0.5}, "views": [lat, ant]}
        r = _save(c, payload, [lat_img, ant_img])
        assert r.status_code == 201, r.text
        ev = r.json()
        assert [v["view"] for v in ev["results"]] == ["anterior", "lateral_derecha"]  # orden de vistas
        assert ev["notes"] == "Control inicial"
        assert ev["patient"]["name"] == "Ana Torres" and ev["age"] == 30
        assert len(ev["results"][0]["tables"]) == 3
        assert "postura" in ev["diagnostic"]["lateral_derecha"]

        # Los archivos se guardaron y se sirven
        res = ev["results"][0]
        for url in (res["image_url"], res["overlay_url"]):
            f = c.get(url)
            assert f.status_code == 200 and f.headers["content-type"] == "image/jpeg"
        assert Image.open(io.BytesIO(c.get(res["image_url"]).content)).size == (res["image_width"],
                                                                               res["image_height"])

        pdf = c.get(ev["report_url"])
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
        assert "inline" in pdf.headers["content-disposition"]
        assert "attachment" in c.get(ev["report_url"], params={"download": True}).headers["content-disposition"]

        xlsx = c.get("/api/evaluations/export.xlsx", params={"patient_id": pid})
        assert xlsx.status_code == 200
        wb = load_workbook(io.BytesIO(xlsx.content))
        assert wb.sheetnames == ["Anterior", "Posterior", "LateralD", "LateralI"]
        assert wb["Anterior"].cell(row=4, column=6).value == "Ana Torres"
        assert wb["Posterior"].max_row == 3  # solo encabezados
        assert wb["Anterior"]["A4"].hyperlink.target.endswith(f"/api/evaluations/{ev['id']}/report.pdf")


def test_server_recomputes_results_from_points():
    """Los resultados salen de los puntos enviados, no de lo que calculó el navegador."""
    with TestClient(app) as c:
        pid = _patient(c)
        ant, img = _analyze(c, "anterior", ANTERIOR)
        ant["landmarks"]["F4"]["y"] += 30  # el evaluador baja el hombro izquierdo
        r = _save(c, {"patient_id": pid, "mode": "markers", "views": [ant]}, [img])
        assert r.status_code == 201, r.text
        assert r.json()["results"][0]["tables"][0]["rows"][0]["direction"] == "Izq."


def test_markerless_evaluation_keeps_height_and_pending_points():
    pose = frontal_pose()
    cal = {"vertex": {"x": 500, "y": 100}, "floor": {"x": 500, "y": 1800}}
    with TestClient(app) as c:
        pid = _patient(c, "Luis Mora")
        r = c.post("/api/analysis/anterior/pose", json={
            "pose": [{"x": q.x, "y": q.y, "visibility": q.visibility} for q in pose],
            "height_cm": 170, "calibration": cal})
        assert r.status_code == 200, r.text
        landmarks = r.json()["markers"]
        img = _png(make_image([], size=(1900, 1000)))
        view = {"view": "anterior", "landmarks": landmarks, "calibration": cal, "estimated_points": ["F7", "F8"],
                "vertical_line": [{"x": 900, "y": 200}, {"x": 900, "y": 1700}]}
        r = _save(c, {"patient_id": pid, "mode": "pose", "height_cm": 170, "views": [view]}, [img])
        assert r.status_code == 201, r.text
        ev = r.json()
        assert abs(ev["results"][0]["scale_ratio"] - 0.1) < 1e-9
        assert ev["results"][0]["estimated_points"] == ["F7", "F8"] and ev["pending_points"] == 2
        assert c.get(f"/api/patients/{pid}").json()["height_cm"] == 170  # la ficha guarda la última estatura

        bad = _save(c, {"patient_id": pid, "mode": "pose", "views": [view]}, [img])  # falta la estatura
        assert bad.status_code == 422
        bad = _save(c, {"patient_id": pid, "mode": "pose", "height_cm": 170, "views": [view, view]}, [img, img])
        assert bad.status_code == 422  # vista repetida
        bad = _save(c, {"patient_id": pid, "mode": "pose", "height_cm": 170, "views": [view]}, [img, img])
        assert bad.status_code == 422  # una foto por vista
        bad = _save(c, {"patient_id": 99999, "mode": "pose", "height_cm": 170, "views": [view]}, [img])
        assert bad.status_code == 404


def test_history_comparison_and_patient_list():
    with TestClient(app) as c:
        pid = _patient(c, "Rosa Paz")
        ant, img = _analyze(c, "anterior", ANTERIOR)
        first = _save(c, {"patient_id": pid, "mode": "markers", "views": [ant]}, [img]).json()
        ant["landmarks"]["F4"]["y"] += 30
        second = _save(c, {"patient_id": pid, "mode": "markers", "views": [ant]}, [img]).json()

        h = c.get(f"/api/patients/{pid}/history", params={"view": "anterior"}).json()
        assert [e["id"] for e in h["evaluations"]] == [first["id"], second["id"]]
        shoulder = h["tables"][0]["rows"][0]
        assert shoulder["values"][1]["direction"] == "Izq."
        assert shoulder["change"] == round(shoulder["values"][1]["value"] - shoulder["values"][0]["value"], 2)
        assert c.get(f"/api/patients/{pid}/history", params={"view": "posterior"}).json()["evaluations"] == []

        listed = next(p for p in c.get("/api/patients").json() if p["id"] == pid)
        assert listed["evaluation_count"] == 2 and listed["last_evaluation_at"].endswith(("Z", "+00:00"))
        assert [e["id"] for e in c.get("/api/evaluations", params={"patient_id": pid}).json()] == [second["id"],
                                                                                                    first["id"]]


def test_delete_evaluation_and_patient_remove_files():
    storage = get_settings().storage_dir
    with TestClient(app) as c:
        pid = _patient(c, "Borrar Prueba")
        ant, img = _analyze(c, "anterior", ANTERIOR)
        e1 = _save(c, {"patient_id": pid, "mode": "markers", "views": [ant]}, [img]).json()
        e2 = _save(c, {"patient_id": pid, "mode": "markers", "views": [ant]}, [img]).json()
        assert (storage / "evaluaciones" / str(e1["id"])).is_dir()

        assert c.delete(f"/api/evaluations/{e1['id']}").status_code == 204
        assert c.get(f"/api/evaluations/{e1['id']}").status_code == 404
        assert not (storage / "evaluaciones" / str(e1["id"])).exists()

        r = c.put(f"/api/patients/{pid}", json={"name": "borrar  prueba", "age": 31, "occupation": " Docente "})
        assert r.json()["age"] == 31 and r.json()["occupation"] == "Docente" and r.json()["name"] == "Borrar Prueba"

        assert c.delete(f"/api/patients/{pid}").status_code == 204
        assert c.get(f"/api/patients/{pid}").status_code == 404
        assert c.get(f"/api/evaluations/{e2['id']}").status_code == 404
        assert not (storage / "evaluaciones" / str(e2["id"])).exists()
